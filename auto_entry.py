from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from playwright.sync_api import (
    sync_playwright,
    TimeoutError as PlaywrightTimeoutError,
)

# =========================
# KONFIGURASI
# =========================

URL = ""
EXCEL_FILE = "data.xlsx"
EXCEL_SHEET = "Sheet1"
STATUS_COLUMN = "STATUS_SITB"

# True = hanya mengisi dan memeriksa form
# False = mengizinkan klik Simpan
DRY_RUN = False

# True = Otomatis klik Simpan dan otomatis 'Yes' tanpa perlu tekan ENTER di terminal
# False = Tetap meminta konfirmasi manual di terminal sebelum Simpan
AUTO_CONFIRM = True

# Jumlah baris yang diproses per sesi (misal: 1, 5, 25, atau None untuk memproses semua yang tersisa)
MAX_ROWS = 5

# Kolom Excel -> ID field teks pada form tambah
TEXT_FIELDS = {
    "dt_tgl_skrining": "dt_tgl_skrining",
    "nik": "nik",
    "berat_badan": "berat_badan",
    "tinggi_badan": "tinggi_badan",
}

# Kolom Excel -> name input dropdown pada form tambah.
# Nama ini harus sesuai hasil inspect_form.py.
DROPDOWN_FIELDS = {
    "tempat_skrining_id_input": "tempat_skrining_id_input",
    "pekerjaan_id_input": "pekerjaan_id_input",
    "riwayat_kontak_tb_id_input": "riwayat_kontak_tb_id_input",
    "risiko_1_id_input": "risiko_1_id_input",
    "risiko_4_id_input": "risiko_4_id_input",
    "risiko_6_id_input": "risiko_6_id_input",
    "risiko_7_id_input": "risiko_7_id_input",
    "risiko_9_id_input": "risiko_9_id_input",
    "cxr_pemeriksaan_id_input": "cxr_pemeriksaan_id_input",
}

# Field yang dapat bersifat kondisional.
OPTIONAL_FIELDS = {
    "berat_badan",
    "tinggi_badan",
    "cxr_alasan",
}


def clean_value(value):
    """Konversi nilai Excel menjadi teks bersih."""
    if value is None or pd.isna(value):
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def get_form_frame(page):
    """Cari iframe yang memuat form tambah SITB (/Skrining/add)."""
    for frame in page.frames:
        if "/Skrining/add" in frame.url:
            return frame

    return None


def get_view_frame(page):
    """Cari iframe yang memuat tampilan hasil simpan SITB (/Skrining/view)."""
    for frame in page.frames:
        if "/Skrining/view" in frame.url:
            return frame

    return None


def get_field_value(frame, field_name):
    """
    Baca nilai dari input HTML atau widget Kendo berdasarkan name atau id.
    Mengembalikan None jika elemen tidak ditemukan di DOM, atau string nilai jika ditemukan.
    """
    try:
        return frame.evaluate("""(name) => {
            const el = document.querySelector(`[name="${name}"]`) || document.getElementById(name);
            if (!el) return null;

            let val = el.value ?? el.getAttribute('value') ?? '';
            if (!val && window.jQuery) {
                const kWidget = $(el).data('kendoComboBox')
                    || $(el).data('kendoDropDownList')
                    || $(el).data('kendoDatePicker')
                    || $(el).data('kendoNumericTextBox');
                if (kWidget && typeof kWidget.value === 'function') {
                    val = kWidget.text ? kWidget.text() : kWidget.value();
                } else {
                    val = $(el).val() || '';
                }
            }
            return String(val || '').trim();
        }""", field_name)
    except Exception:
        return ""


def wait_for_form_ready(frame, timeout=20000):
    """
    Tunggu hingga dokumen iframe /Skrining/add benar-benar selesai dimuat (readyState === 'complete')
    dan widget Kendo DatePicker pada #dt_tgl_skrining sudah siap.
    """
    frame.wait_for_function(
        """() => document.readyState === 'complete'
            && !!document.getElementById('dt_tgl_skrining')
            && !!document.getElementById('nik')
            && (!window.jQuery || !!$('#dt_tgl_skrining').data('kendoDatePicker'))""",
        timeout=timeout,
    )
    frame.locator("#dt_tgl_skrining").wait_for(state="visible", timeout=10000)
    frame.locator("#nik").wait_for(state="visible", timeout=10000)


def close_dialog_modal(page, wait_seconds=3):
    """
    Menutup pop-up modal Kendo Window (#dialog) baik saat menampilkan /Skrining/view (setelah simpan)
    maupun saat masih di /Skrining/add (ketika terjadi GAGAL_NIK), lalu menunggu wait_seconds detik.
    """
    try:
        dialog = page.locator("#dialog")
        is_modal_open = False
        try:
            is_modal_open = dialog.count() > 0 and dialog.first.is_visible()
        except Exception:
            is_modal_open = False

        if not is_modal_open:
            # Cek juga apakah ada container .k-window yang sedang visible
            k_win = page.locator(".k-window:visible")
            is_modal_open = k_win.count() > 0

        if is_modal_open:
            print(f"-> Menutup pop-up modal (#dialog) dan menunggu {wait_seconds} detik...")

            # 1. Coba klik tombol Close (X) pada titlebar Kendo Window sesuai struktur HTML SITB
            close_selectors = [
                ".k-window:visible .k-window-actions a.k-window-action:has(.k-i-close)",
                ".k-window:visible .k-window-titlebar .k-i-close",
                ".k-window:visible .k-window-actions a.k-window-action",
            ]
            clicked = False
            for sel in close_selectors:
                try:
                    btn = page.locator(sel)
                    if btn.count() > 0 and btn.first.is_visible():
                        btn.first.click(timeout=3000)
                        clicked = True
                        break
                except Exception:
                    continue

            page.wait_for_timeout(500)

            # 2. Fallback via API Kendo Window jika modal masih terlihat setelah klik tombol X
            try:
                still_open = dialog.count() > 0 and dialog.first.is_visible()
                if not clicked or still_open:
                    page.evaluate("""() => {
                        if (window.jQuery) {
                            const dlg = $('#dialog').data('kendoWindow');
                            if (dlg) dlg.close();
                            $('.k-window-content:visible').each(function() {
                                const w = $(this).data('kendoWindow');
                                if (w) w.close();
                            });
                        }
                    }""")
            except Exception:
                pass

            # 3. Tunggu hingga modal #dialog benar-benar tertutup (hidden)
            try:
                dialog.first.wait_for(state="hidden", timeout=10000)
            except Exception:
                pass

            # 4. Tunggu sesuai jeda yang diminta (default 3 detik) sebelum aksi berikutnya
            page.wait_for_timeout(int(wait_seconds * 1000))
            print("-> Modal berhasil ditutup.")
    except Exception as e:
        print(f"-> Info penutupan modal: {e}")
        page.wait_for_timeout(int(wait_seconds * 1000))


def reset_add_form(page, frame):
    """
    Saat terjadi GAGAL_NIK atau kendala input, tutup modal terlebih dahulu
    dan tunggu 3 detik agar baris berikutnya selalu dimulai dari klik Tambah yang baru.
    """
    handle_confirmation_dialog(page, frame)
    close_dialog_modal(page, wait_seconds=3)


def open_add_form(page):
    """
    Pastikan tidak ada modal (#dialog) yang masih terbuka dari baris sebelumnya,
    lalu klik tombol Tambah (#btnAdd_ta_skrining) dan tunggu form /Skrining/add siap.
    """
    # Pastikan modal sebelumnya (misal /Skrining/view?id=... atau sisa form lama) sudah ditutup
    dialog = page.locator("#dialog")
    try:
        if dialog.count() > 0 and dialog.first.is_visible():
            close_dialog_modal(page, wait_seconds=3)
    except Exception:
        pass

    button = page.locator("#btnAdd_ta_skrining")
    button.wait_for(state="visible", timeout=15000)
    button.click()

    try:
        page.wait_for_function(
            """() => Array.from(document.querySelectorAll("iframe"))
                .some(f => (f.src || "").includes("/Skrining/add") && !!(f.offsetWidth || f.offsetHeight))""",
            timeout=15000,
        )
    except PlaywrightTimeoutError:
        pass

    page.wait_for_timeout(1500)
    frame = get_form_frame(page)

    if frame is None:
        raise RuntimeError(
            "Form iframe /Skrining/add tidak ditemukan. "
            "Pastikan tombol Tambah berhasil membuka form."
        )

    wait_for_form_ready(frame, timeout=20000)
    page.wait_for_timeout(1000)

    print("Form ditemukan & siap:", frame.url)
    return frame


def is_valid_nik_format(nik_str):
    """Periksa apakah NIK terdiri dari 16 digit angka dengan kode tanggal lahir yang wajar."""
    if not nik_str:
        return False, "NIK kosong"

    nik_clean = str(nik_str).strip()
    if len(nik_clean) != 16 or not nik_clean.isdigit():
        return False, f"Panjang NIK {len(nik_clean)} digit (wajib 16 digit angka)"

    dd = int(nik_clean[6:8])
    mm = int(nik_clean[8:10])
    real_dd = dd - 40 if dd > 40 else dd

    if real_dd < 1 or real_dd > 31 or mm < 1 or mm > 12:
        return False, f"Kode tanggal/bulan lahir pada NIK tidak valid (dd={dd}, mm={mm})"

    return True, "OK"


def fill_text_field(frame, name, value):
    if value is None or value == "":
        print(f"SKIP field kosong: {name}")
        return

    field = frame.locator(f'[name="{name}"], #{name}')

    if field.count() == 0:
        print(f"SKIP field tidak ditemukan sama sekali: {name}")
        return

    # Penanganan khusus untuk NIK: klik, isi, dan tekan Enter sekali layaknya input manual
    if name == "nik":
        val_str = str(value).strip()
        try:
            target = field.first
            target.wait_for(state="visible", timeout=10000)
            target.click(timeout=3000)
            target.fill(val_str)
            frame.page.wait_for_timeout(300)
            target.press("Enter")
            print(f"OK isi & Enter: {name}")
            return
        except Exception:
            # Fallback via JS jika elemen tidak bisa diklik langsung
            try:
                field.first.evaluate("""(el, val) => {
                    el.value = val;
                    el.dispatchEvent(new Event('input', { bubbles: true }));
                    el.dispatchEvent(new Event('change', { bubbles: true }));
                    if (window.jQuery) {
                        $(el).trigger(jQuery.Event('keydown', { which: 13, keyCode: 13 }));
                        $(el).trigger(jQuery.Event('keypress', { which: 13, keyCode: 13 }));
                        $(el).trigger(jQuery.Event('keyup', { which: 13, keyCode: 13 }));
                        $(el).trigger('change');
                    }
                }""", val_str)
                print(f"OK isi (JS Fallback): {name}")
                return
            except Exception as e:
                print(f"GAGAL isi {name}: {e}")
                return

    # Penanganan untuk field teks lainnya (dt_tgl_skrining, berat_badan, tinggi_badan)
    try:
        field.first.evaluate("""(el, rawValue) => {
            let val = String(rawValue || '').split(' ')[0]; // Ambil nilai/tanggal tanpa jam

            // 1. Paksa hapus readonly agar kita bisa injeksi nilai manual jika Kendo gagal
            el.removeAttribute('readonly');

            if (window.jQuery) {
                let kendoNumeric = $(el).data('kendoNumericTextBox');
                if (kendoNumeric) {
                    kendoNumeric.value(parseFloat(val));
                    kendoNumeric.trigger('change');
                    return;
                }

                let kendoDate = $(el).data('kendoDatePicker');
                if (kendoDate) {
                    // Kendo DatePicker membutuhkan JS Date Object, mendukung pemisah '-' maupun '/'
                    let parts = val.split(/[-/]/);
                    let dateObj = null;
                    if (parts.length === 3) {
                        if (parts[0].length === 4) {
                            // Format: YYYY-MM-DD
                            dateObj = new Date(parseInt(parts[0]), parseInt(parts[1]) - 1, parseInt(parts[2]));
                        } else {
                            // Format: DD-MM-YYYY atau DD/MM/YYYY
                            dateObj = new Date(parseInt(parts[2]), parseInt(parts[1]) - 1, parseInt(parts[0]));
                        }
                    } else {
                        dateObj = new Date(val);
                    }

                    // Biarkan Kendo memformat teks sesuai options.format bawaan form SITB
                    kendoDate.value(dateObj);
                    kendoDate.trigger('change');

                    // Hanya isi manual jika kendoDate.value() tidak mengisi el.value
                    if (!el.value) {
                        el.value = val;
                        $(el).trigger('change');
                    }

                    // Sinkronkan ke hidden input (tgl_skrining) jika ada
                    let hiddenId = el.id.replace('dt_', '');
                    let hiddenEl = document.getElementById(hiddenId);
                    if (hiddenEl) {
                        hiddenEl.value = el.value || val;
                    }
                    return;
                }
            }

            // Fallback untuk text field standar
            el.value = val;
            el.dispatchEvent(new Event('input', { bubbles: true }));
            el.dispatchEvent(new Event('change', { bubbles: true }));
        }""", value)
        print(f"OK isi: {name}")
    except Exception as e:
        print(f"GAGAL isi {name}: {e}")


def verify_nik_by_nama_peserta(frame, nik_val=None):
    """
    Mengecek apakah NIK valid dengan memeriksa field 'nama_peserta'.
    Mengantisipasi trafik server SITB yang padat dengan 3 tahap pengecekan
    (cek awal 4 detik + 2 kali cek ulang @ 4 detik = total maksimal 12 detik).
    Jika nama_peserta muncul lebih cepat, langsung lanjut tanpa menunggu sisa waktu.
    """
    initial_check = get_field_value(frame, "nama_peserta")
    if initial_check is None:
        # Jika elemen nama_peserta tidak ada di DOM sama sekali, jangan gagalkan baris
        frame.page.wait_for_timeout(2000)
        return True, "(Field nama_peserta tidak ditemukan di DOM)"

    max_cycles = 3       # 1 cek awal + 2 kali cek ulang
    wait_per_cycle = 4   # 4 detik per tahap (total maksimal 12 detik)

    for cycle in range(1, max_cycles + 1):
        for _ in range(wait_per_cycle):
            frame.page.wait_for_timeout(1000)
            nama_val = get_field_value(frame, "nama_peserta")
            if nama_val:
                # Beri jeda 1 detik agar field lain (provinsi/kab/jk/umur/gejala) selesai ter-render dari respons NIK
                frame.page.wait_for_timeout(1000)
                return True, nama_val

        if cycle < max_cycles:
            print(
                f"-> [Cek NIK {cycle}/{max_cycles}] 'nama_peserta' belum tampil ({cycle * wait_per_cycle} detik), "
                f"menunggu {wait_per_cycle} detik lagi (percobaan ke-{cycle + 1})..."
            )
            # Pastikan nilai NIK masih ada di kotak input (jika sempat terhapus, isi ulang & tekan Enter)
            if cycle == 1 or cycle == 2:
                try:
                    curr_nik = get_field_value(frame, "nik")
                    nik_el = frame.locator('#nik, [name="nik"]').first
                    if nik_val and curr_nik != str(nik_val).strip():
                        nik_el.click(timeout=2000)
                        nik_el.fill(str(nik_val).strip())
                        frame.page.wait_for_timeout(300)
                        nik_el.press("Enter")
                    elif cycle == 2:
                        nik_el.press("Enter")
                except Exception:
                    pass

    return False, ""


def fill_wilayah_ktp_if_empty(frame, row):
    """
    Jika provinsi_ktp_id_input atau kabupaten_ktp_id_input masih kosong (null)
    setelah generate NIK, ambil nilainya dari Excel dan isikan ke form.
    """
    # 1. Cek Provinsi KTP
    curr_prov = get_field_value(frame, "provinsi_ktp_id_input")
    if curr_prov is not None:
        if not curr_prov:
            excel_prov = clean_value(row.get("provinsi_ktp_id_input"))
            if excel_prov:
                print(f"-> provinsi_ktp_id_input kosong di web, mengambil dari Excel: '{excel_prov}'")
                select_kendo_dropdown(frame, "provinsi_ktp_id_input", excel_prov)
                # Tunggu cascading dropdown kabupaten_ktp_id_input selesai memuat daftar kab/kota dari server
                frame.page.wait_for_timeout(1500)
        else:
            print(f"OK provinsi_ktp_id_input sudah terisi: '{curr_prov}'")

    # 2. Cek Kabupaten/Kota KTP
    curr_kab = get_field_value(frame, "kabupaten_ktp_id_input")
    if curr_kab is not None:
        if not curr_kab:
            excel_kab = clean_value(row.get("kabupaten_ktp_id_input"))
            if excel_kab:
                print(f"-> kabupaten_ktp_id_input kosong di web, mengambil dari Excel: '{excel_kab}'")
                select_kendo_dropdown(frame, "kabupaten_ktp_id_input", excel_kab)
                frame.page.wait_for_timeout(1000)
        else:
            print(f"OK kabupaten_ktp_id_input sudah terisi: '{curr_kab}'")


def select_kendo_dropdown(frame, name, value):
    if value is None or value == "":
        print(f"SKIP dropdown kosong: {name}")
        return

    # Cari elemen visible (biasanya yang ada _input di namanya)
    dropdown = frame.locator(f'[name="{name}"]:visible')

    if dropdown.count() > 0:
        try:
            dropdown.click(timeout=3000)
            dropdown.press("Control+a")
            dropdown.press("Backspace")
            dropdown.fill(value)
            # Tunggu 1 detik agar popup daftar opsi Kendo muncul (JANGAN tekan Enter dulu agar popup tidak tertutup)
            frame.page.wait_for_timeout(1000)
        except Exception:
            pass  # Lanjut ke pencarian opsi / JS Bypass jika gagal klik
    else:
        # Jika benar-benar tersembunyi
        dropdown = frame.locator(f'[name="{name}"]')
        if dropdown.count() == 0:
            print(f"SKIP dropdown tidak ditemukan: {name}")
            return

    page = frame.page
    # Prioritaskan kecocokan persis (exact=True) terlebih dahulu agar 'Tidak' tidak salah klik 'Tidak Diketahui'
    candidates = [
        frame.get_by_role("option", name=value, exact=True),
        page.get_by_role("option", name=value, exact=True),
        frame.get_by_role("option", name=value),
        page.get_by_role("option", name=value),
    ]

    selected = False
    for candidate in candidates:
        try:
            if candidate.count() > 0:
                candidate.first.wait_for(state="visible", timeout=1500)
                candidate.first.click()
                selected = True
                break
        except PlaywrightTimeoutError:
            continue
        except Exception:
            continue

    # ==========================================
    # BYPASS JAVASCRIPT FINAL UNTUK KENDO UI
    # ==========================================
    if not selected:
        print(f"WARNING: Gagal klik opsi '{name}'. Menggunakan Bypass JS...")

        # Contoh: dari "gejala_1_1_id_input" menjadi "gejala_1_1_id"
        base_name = name.replace("_input", "") if name.endswith("_input") else name

        try:
            js_success = frame.evaluate("""(params) => {
                let { baseName, inputName, val } = params;

                // 1. Ambil elemen asli (hidden) dan elemen visual (input teks)
                let baseEl = document.getElementById(baseName);
                let visualEl = document.querySelector(`input[name="${inputName}"]`);

                if (window.jQuery && baseEl) {
                    let widget = $(baseEl).data('kendoComboBox') || $(baseEl).data('kendoDropDownList');
                    if (widget) {
                        let data = (widget.dataSource && widget.dataSource.data()) || [];
                        let found = false;
                        let valLower = String(val).toLowerCase().trim();

                        let textField = (widget.options && widget.options.dataTextField) || 'text';
                        let valueField = (widget.options && widget.options.dataValueField) || 'value';

                        // Tahap A: Cari yang cocok persis (exact match) terlebih dahulu
                        for (let i = 0; i < data.length; i++) {
                            let itemText = String(data[i][textField] || data[i].text || '').toLowerCase().trim();
                            if (itemText === valLower) {
                                let actualVal = data[i][valueField] !== undefined ? data[i][valueField] : data[i].value;
                                widget.value(actualVal);
                                if (widget.text) widget.text(data[i][textField] || data[i].text);
                                widget.trigger('change');
                                if (baseEl) {
                                    baseEl.value = actualVal;
                                    baseEl.dispatchEvent(new Event('change', { bubbles: true }));
                                }
                                found = true;
                                break;
                            }
                        }

                        // Tahap B: Jika tidak ada exact match, cari yang mengandung teks tersebut (includes)
                        if (!found) {
                            for (let i = 0; i < data.length; i++) {
                                let itemText = String(data[i][textField] || data[i].text || '').toLowerCase().trim();
                                if (itemText && itemText.includes(valLower)) {
                                    let actualVal = data[i][valueField] !== undefined ? data[i][valueField] : data[i].value;
                                    widget.value(actualVal);
                                    if (widget.text) widget.text(data[i][textField] || data[i].text);
                                    widget.trigger('change');
                                    if (baseEl) {
                                        baseEl.value = actualVal;
                                        baseEl.dispatchEvent(new Event('change', { bubbles: true }));
                                    }
                                    found = true;
                                    break;
                                }
                            }
                        }

                        // Jika data belum ter-load dari server, paksa nilai masuk
                        if (!found) {
                            if (widget.text) widget.text(val);
                            widget.value(val);
                            widget.trigger('change');
                            if (baseEl) {
                                baseEl.value = val;
                                baseEl.dispatchEvent(new Event('change', { bubbles: true }));
                            }
                        }
                        return true;
                    }
                }

                // 2. INJEKSI KE TAMPILAN HTML
                if (visualEl) {
                    visualEl.value = val;
                    visualEl.dispatchEvent(new Event('input', { bubbles: true }));
                    visualEl.dispatchEvent(new Event('change', { bubbles: true }));
                }

                // 3. INJEKSI KE ELEMEN HIDDEN UNTUK DISUBMIT
                if (baseEl) {
                    baseEl.value = val;
                    baseEl.dispatchEvent(new Event('change', { bubbles: true }));
                }

                return true;
            }""", {"baseName": base_name, "inputName": name, "val": str(value)})

            if not js_success:
                print(f"GAGAL: Gagal injeksi JS untuk '{name}'.")
            else:
                print(f"OK pilih dropdown (via JS Bypass): {name}")
                return
        except Exception as err:
            print(f"GAGAL: Eksepsi injeksi JS untuk '{name}': {err}")

    print(f"OK pilih dropdown: {name}")


def fill_symptoms_dynamic(frame, row):
    """
    Mengisi form skrining gejala & tanda secara adaptif.
    Mendeteksi apakah pasien masuk kategori ANAK (gejala_1_*) atau DEWASA (gejala_2_*),
    lalu memetakan data dari kolom Excel secara otomatis.
    """
    frame.page.wait_for_timeout(800)

    # Deteksi form Anak: periksa apakah kontainer gejala_1_1_id terlihat (display != none)
    anak_container = frame.locator("#form_item_gejala_1_1_id")
    is_anak = False
    if anak_container.count() > 0:
        try:
            is_anak = anak_container.is_visible()
        except Exception:
            is_anak = False

    if is_anak:
        print("-> Terdeteksi form skrining ANAK (< 15 tahun). Memetakan ke kelompok gejala_1_*...")
        batuk_val = clean_value(row.get("gejala_1_1_id_input") or row.get("gejala_2_1_id_input"))
        bb_val = clean_value(row.get("gejala_1_3_id_input") or row.get("gejala_2_3_id_input"))
        demam_val = clean_value(row.get("gejala_1_4_id_input") or row.get("gejala_2_4_id_input"))
        lesu_val = clean_value(row.get("gejala_1_5_id_input") or row.get("gejala_2_5_id_input"))

        gejala_map = {
            "gejala_1_1_id_input": batuk_val,
            "gejala_1_3_id_input": bb_val,
            "gejala_1_4_id_input": demam_val,
            "gejala_1_5_id_input": lesu_val,
        }
    else:
        print("-> Terdeteksi form skrining DEWASA (>= 15 tahun). Memetakan ke kelompok gejala_2_*...")
        gejala_map = {
            "gejala_2_1_id_input": clean_value(row.get("gejala_2_1_id_input")),
            "gejala_2_3_id_input": clean_value(row.get("gejala_2_3_id_input")),
            "gejala_2_4_id_input": clean_value(row.get("gejala_2_4_id_input")),
            "gejala_2_5_id_input": clean_value(row.get("gejala_2_5_id_input")),
        }

    # Pembesaran kelenjar getah bening (gejala_6_id_input): hanya jika terlihat di form
    g6_val = clean_value(row.get("gejala_6_id_input"))
    if g6_val:
        g6_container = frame.locator("#form_item_gejala_6_id")
        if g6_container.count() > 0 and g6_container.is_visible():
            gejala_map["gejala_6_id_input"] = g6_val

    for input_name, val in gejala_map.items():
        if val:
            select_kendo_dropdown(frame, input_name, val)


def fill_optional_cxr_reason(frame, row):
    value = clean_value(row.get("cxr_alasan"))

    if not value:
        return

    field = frame.locator('[name="cxr_alasan"]:visible')

    if field.count() == 0:
        print(
            "PERINGATAN: cxr_alasan ada di Excel, "
            "tetapi field tidak terlihat pada form."
        )
        return

    field.fill(value)
    print("OK isi: cxr_alasan")


def handle_confirmation_dialog(page, frame):
    """
    Otomatis mencari dan mengklik tombol konfirmasi ('Ya' / 'Yes' / 'OK' / 'Setuju')
    jika muncul popup konfirmasi atau notifikasi SITB.
    """
    page.wait_for_timeout(1000)

    selectors = [
        "button:has-text('Ya')",
        "button:has-text('Yes')",
        "button:has-text('OK')",
        "button:has-text('Setuju')",
        ".swal2-confirm",
        ".bootbox-accept",
        ".k-dialog-button-layout .k-primary",
    ]

    for target in [page, frame]:
        if target is None:
            continue
        for selector in selectors:
            try:
                btn = target.locator(selector)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click()
                    print(f"-> Otomatis klik konfirmasi dialog: '{selector}'")
                    page.wait_for_timeout(1200)
                    return True
            except Exception:
                continue
    return False


def wait_after_save(page, frame):
    """
    Menunggu proses penyimpanan selesai setelah klik tombol #save:
    - Mendeteksi peralihan iframe modal (#dialog) dari /Skrining/add ke /Skrining/view?id=...
    - Atau memeriksa apakah muncul pesan validasi error (.k-invalid-msg)
    """
    page.wait_for_timeout(1500)
    handle_confirmation_dialog(page, frame)

    for _ in range(24):  # 24 x 500ms = maksimal 12 detik menunggu respons server
        # 1. Cek apakah iframe di dalam modal #dialog sudah beralih ke /Skrining/view?id=...
        view_frame = get_view_frame(page)
        if view_frame is not None:
            print(f"-> Simpan berhasil! Pop-up View Data tampil: {view_frame.url}")
            return True, []

        # 2. Cek juga via atribut src pada elemen iframe.k-content-frame di #dialog
        try:
            iframe_src = page.locator("#dialog iframe.k-content-frame").get_attribute("src") or ""
            if "/Skrining/view" in iframe_src:
                print(f"-> Simpan berhasil! Pop-up View Data tampil: {iframe_src}")
                return True, []
        except Exception:
            pass

        # 3. Cek jika iframe /Skrining/add sudah tidak ada
        current_frame = get_form_frame(page)
        if current_frame is None:
            return True, []

        # 4. Cek apakah muncul pesan validasi error (.k-invalid-msg) di dalam form /Skrining/add
        try:
            invalid_msgs = current_frame.evaluate("""() => {
                return Array.from(document.querySelectorAll('.k-invalid-msg'))
                    .filter(el => (el.offsetWidth || el.offsetHeight) && el.innerText.trim() !== '')
                    .map(el => el.innerText.trim());
            }""")
            if invalid_msgs:
                return False, invalid_msgs
        except Exception:
            pass

        page.wait_for_timeout(500)

    return True, []


def mark_excel_status(file_path, sheet_name, row_idx, status="SUKSES"):
    """
    Catat status ke file Excel menggunakan openpyxl agar aman, instan,
    dan tidak merusak format sel lainnya.
    row_idx: 0-indexed dari pandas DataFrame (baris di Excel = row_idx + 2)
    """
    excel_row = row_idx + 2
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

        headers = [str(cell.value) if cell.value is not None else "" for cell in ws[1]]

        if STATUS_COLUMN not in headers:
            status_col = len(headers) + 1
            ws.cell(row=1, column=status_col, value=STATUS_COLUMN)
            headers.append(STATUS_COLUMN)
        else:
            status_col = headers.index(STATUS_COLUMN) + 1

        if "WAKTU_ENTRY" not in headers:
            time_col = len(headers) + 1
            ws.cell(row=1, column=time_col, value="WAKTU_ENTRY")
            headers.append("WAKTU_ENTRY")
        else:
            time_col = headers.index("WAKTU_ENTRY") + 1

        ws.cell(row=excel_row, column=status_col, value=status)
        ws.cell(row=excel_row, column=time_col, value=timestamp)

        wb.save(file_path)
        print(f"-> [Excel] Baris {excel_row} berhasil ditandai '{status}' ({timestamp})")
    except PermissionError:
        print(
            f"-> [PERINGATAN] File Excel '{file_path}' sedang dibuka di aplikasi lain. "
            "Tutup file Excel agar status tersimpan ke file."
        )
    except Exception as e:
        print(f"-> [PERINGATAN] Gagal mencatat status ke Excel: {e}")


def verify_text_fields(frame, row):
    """Verifikasi nilai teks dengan toleransi format tanggal dan angka."""
    print("\nVerifikasi field teks:")

    for column, field_id in TEXT_FIELDS.items():
        expected = clean_value(row.get(column))

        if not expected:
            continue

        field = frame.locator(f"#{field_id}")

        if field.count() == 0:
            print(f"PERIKSA: #{field_id} tidak ditemukan di halaman.")
            continue

        actual = field.input_value()

        # Jika dt_tgl_skrining sempat kosong akibat proses loading form, isi ulang sekali sebelum verifikasi
        if column == "dt_tgl_skrining" and not actual:
            fill_text_field(frame, field_id, expected)
            actual = field.input_value()

        matches = False

        # 1. Logika khusus perbandingan Tanggal
        if column == "dt_tgl_skrining":
            try:
                exp_clean = expected.replace("/", "-")
                act_clean = actual.replace("/", "-")

                exp_date = pd.to_datetime(exp_clean, dayfirst=True).date()
                act_date = pd.to_datetime(act_clean, dayfirst=True).date()
                matches = exp_date == act_date
            except Exception:
                matches = actual == expected

        # 2. Logika khusus perbandingan Angka (Berat/Tinggi)
        elif column in ["berat_badan", "tinggi_badan"]:
            try:
                matches = float(actual) == float(expected)
            except ValueError:
                matches = actual == expected

        # 3. Logika untuk teks biasa (NIK)
        else:
            matches = actual == expected

        print(
            f"  {field_id}: {'SESUAI' if matches else 'TIDAK SESUAI'} "
            f"(Excel: '{expected}' | Web: '{actual}')"
        )

        if not matches:
            raise ValueError(
                f"Verifikasi gagal untuk field '{field_id}'. "
                f"Diinginkan: '{expected}', tetapi web berisi: '{actual}'."
            )


def ensure_excel_writable(file_path):
    """Pastikan file Excel tidak sedang dikunci (terbuka di Microsoft Excel) sebelum memulai otomasi."""
    while True:
        try:
            with open(file_path, "a"):
                pass
            break
        except PermissionError:
            input(
                f"\n[PERINGATAN] File '{file_path}' sedang terbuka di Microsoft Excel!\n"
                "Mohon SIMPAN (Ctrl+S) dan TUTUP aplikasi Microsoft Excel terlebih dahulu, "
                "lalu tekan ENTER untuk melanjutkan..."
            )


def main():
    file_path = Path(EXCEL_FILE)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Excel tidak ditemukan: {file_path.resolve()}"
        )

    ensure_excel_writable(file_path)

    # Baca NIK sebagai teks agar angka nol di depan tidak hilang.
    df = pd.read_excel(
        file_path,
        sheet_name=EXCEL_SHEET,
        dtype={"nik": str},
    )

    # Format kolom dt_tgl_skrining menjadi string tanggal DD-MM-YYYY
    if "dt_tgl_skrining" in df.columns:
        df["dt_tgl_skrining"] = pd.to_datetime(
            df["dt_tgl_skrining"], dayfirst=True, errors="coerce"
        ).dt.strftime("%d-%m-%Y")

    if df.empty:
        raise ValueError("Excel tidak memiliki data.")

    required_columns = {
        "dt_tgl_skrining",
        "nik",
        *DROPDOWN_FIELDS.keys(),
    }

    missing = required_columns - set(df.columns)

    if missing:
        raise ValueError(
            "Kolom Excel belum lengkap: "
            + ", ".join(sorted(missing))
        )

    if not any(col in df.columns for col in ["gejala_2_1_id_input", "gejala_1_1_id_input"]):
        raise ValueError(
            "Kolom gejala belum ditemukan di Excel (butuh minimal 'gejala_2_1_id_input' atau 'gejala_1_1_id_input')."
        )

    if STATUS_COLUMN not in df.columns:
        df[STATUS_COLUMN] = None

    status_series = df[STATUS_COLUMN].astype(str).str.strip().str.upper()
    sukses_count = int((status_series == "SUKSES").sum())
    gagal_count = int((status_series == "GAGAL_NIK").sum())

    # Filter hanya baris yang belum SUKSES dan belum GAGAL_NIK
    pending_mask = ~status_series.isin(["SUKSES", "GAGAL_NIK"])
    pending_df = df[pending_mask]

    if pending_df.empty:
        print("\n[INFO] SEMUA DATA DI EXCEL SUDAH DIPROSES (SUKSES / GAGAL_NIK).")
        print("Tidak ada antrean data baru yang perlu di-entry.")
        return

    print("\nStatus Data Excel:")
    print(f"- Total baris di Excel     : {len(df)}")
    print(f"- Sudah berstatus SUKSES   : {sukses_count}")
    print(f"- Berstatus GAGAL_NIK      : {gagal_count}")
    print(f"- Belum diproses (pending) : {len(pending_df)}")

    if MAX_ROWS is not None:
        records = pending_df.head(MAX_ROWS)
    else:
        records = pending_df

    total_records = len(records)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context()
        page = context.new_page()

        # Otomatis klik OK/Accept jika muncul dialog bawaan browser (confirm/alert)
        page.on("dialog", lambda dialog: dialog.accept())

        try:
            page.goto(URL, wait_until="domcontentloaded")

            input(
                "Login SITB dan buka menu Data Skrining, "
                "lalu tekan ENTER..."
            )

            for current_idx, (index, row) in enumerate(records.iterrows(), start=1):
                excel_row_num = index + 2
                nik_val = clean_value(row.get("nik"))
                print(f"\n[{current_idx}/{total_records}] Memproses baris Excel: {excel_row_num} (NIK: {nik_val})")

                # 1. Validasi awal format 16 digit & kode tanggal lahir pada NIK
                valid_fmt, fmt_reason = is_valid_nik_format(nik_val)
                if not valid_fmt:
                    print(
                        f"SKIP Baris {excel_row_num}: Format NIK tidak sesuai ({fmt_reason}) "
                        "-> Ditandai GAGAL_NIK"
                    )
                    mark_excel_status(EXCEL_FILE, EXCEL_SHEET, index, status="GAGAL_NIK")
                    continue

                frame = None
                try:
                    frame = open_add_form(page)

                    # 2. Isi Tanggal Skrining & NIK terlebih dahulu
                    fill_text_field(frame, "dt_tgl_skrining", clean_value(row.get("dt_tgl_skrining")))
                    fill_text_field(frame, "nik", nik_val)

                    # 3. Tunggu dan cek kevalidan NIK menggunakan field nama_peserta
                    #    SEBELUM mengisi field lainnya agar tidak terjadi tabrakan proses loading
                    nik_ok, nama_peserta = verify_nik_by_nama_peserta(frame, nik_val)
                    if not nik_ok:
                        print(
                            f"SKIP Baris {excel_row_num}: Field 'nama_peserta' tidak keluar / kosong "
                            f"untuk NIK {nik_val} setelah 12 detik -> Ditandai GAGAL_NIK"
                        )
                        mark_excel_status(EXCEL_FILE, EXCEL_SHEET, index, status="GAGAL_NIK")
                        reset_add_form(page, frame)
                        continue

                    print(f"OK NIK Valid -> Nama Peserta: {nama_peserta}")

                    # 4. Isi Berat Badan & Tinggi Badan (setelah loading NIK selesai)
                    fill_text_field(frame, "berat_badan", clean_value(row.get("berat_badan")))
                    fill_text_field(frame, "tinggi_badan", clean_value(row.get("tinggi_badan")))

                    # 5. Isi dropdown utama
                    for column, input_name in DROPDOWN_FIELDS.items():
                        select_kendo_dropdown(
                            frame,
                            input_name,
                            clean_value(row.get(column)),
                        )

                    # 6. Isi dropdown gejala adaptif (otomatis mendeteksi form Anak atau Dewasa)
                    fill_symptoms_dynamic(frame, row)

                    # 7. Isi alasan CXR (jika ada)
                    fill_optional_cxr_reason(frame, row)

                    # 8. Pengecekan & pengisian Provinsi / Kabupaten KTP ditaruh PALING AKHIR
                    #    agar tidak bentrok dengan proses auto-load wilayah dari NIK
                    fill_wilayah_ktp_if_empty(frame, row)

                    # 9. Verifikasi akhir field teks sebelum Simpan
                    verify_text_fields(frame, row)

                    print(f"\nPengisian baris {excel_row_num} selesai. Periksa form di browser.")

                    if DRY_RUN:
                        print(
                            "DRY_RUN=True: tidak ada tombol Simpan yang diklik."
                        )
                        if current_idx < total_records:
                            choice = input(
                                "Tekan ENTER untuk lanjut ke baris berikutnya, "
                                "atau ketik 'q' untuk keluar: "
                            ).strip().lower()
                            if choice == "q":
                                break
                            reset_add_form(page, frame)
                        else:
                            input("Semua baris selesai diperiksa. Tekan ENTER untuk menutup browser...")
                        continue

                    # Mode Simpan (DRY_RUN=False)
                    save_button = frame.locator("#save")

                    if save_button.count() != 1:
                        raise ValueError(
                            "Tombol #save tidak ditemukan secara unik "
                            "di iframe form."
                        )

                    if not save_button.is_visible() or not save_button.is_enabled():
                        raise ValueError("Tombol Simpan belum siap digunakan.")

                    if not AUTO_CONFIRM:
                        choice = input(
                            f"PERIKSA DATA DI BROWSER (Baris {excel_row_num}).\n"
                            "Tekan ENTER untuk Simpan baris ini, "
                            "atau ketik 'q' untuk membatalkan proses: "
                        ).strip().lower()

                        if choice == "q":
                            print("Proses dibatalkan oleh pengguna.")
                            break

                    print("Mengklik tombol Simpan...")
                    save_button.click()

                    # Tangani otomatis jika muncul popup konfirmasi "Ya / Yes / OK"
                    handle_confirmation_dialog(page, frame)

                    print(
                        "Tombol Simpan telah diproses. "
                        "Menunggu respons penyimpanan dari SITB..."
                    )
                    save_ok, invalid_msgs = wait_after_save(page, frame)

                    if not save_ok and invalid_msgs:
                        print(
                            f"SKIP Baris {excel_row_num}: Validasi SITB menolak simpan "
                            f"({'; '.join(invalid_msgs)}) -> Ditandai GAGAL_NIK"
                        )
                        mark_excel_status(EXCEL_FILE, EXCEL_SHEET, index, status="GAGAL_NIK")
                        reset_add_form(page, frame)
                        continue

                    # Catat status SUKSES ke file Excel agar tidak diulang lagi
                    mark_excel_status(EXCEL_FILE, EXCEL_SHEET, index, status="SUKSES")

                    # Tutup pop-up modal View Data (/Skrining/view?id=...) dan tunggu 3 detik
                    close_dialog_modal(page, wait_seconds=3)

                    if current_idx < total_records:
                        if not AUTO_CONFIRM:
                            next_choice = input(
                                "Tekan ENTER untuk lanjut ke baris berikutnya, "
                                "atau ketik 'q' untuk berhenti: "
                            ).strip().lower()
                            if next_choice == "q":
                                break
                        else:
                            print("Otomatis melanjutkan ke baris berikutnya...")
                    else:
                        print("\nSemua baris selesai diproses.")
                        if not AUTO_CONFIRM:
                            input("Tekan ENTER untuk menutup browser...")

                except Exception as row_err:
                    print(
                        f"\n[SKIP] Kendala pada Baris {excel_row_num} (NIK: {nik_val}): {row_err}\n"
                        "-> Baris ditandai GAGAL_NIK dan otomatis lanjut ke baris berikutnya..."
                    )
                    mark_excel_status(EXCEL_FILE, EXCEL_SHEET, index, status="GAGAL_NIK")
                    reset_add_form(page, frame)
                    continue

        except Exception as exc:
            print(f"\nGAGAL: {exc}")
            print(
                "Proses dihentikan agar tidak meneruskan "
                "pengisian dengan kondisi yang belum benar."
            )
            input("Periksa browser, lalu tekan ENTER untuk menutup...")

        finally:
            browser.close()


if __name__ == "__main__":
    main()